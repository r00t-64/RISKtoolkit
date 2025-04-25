import pandas as pd
import numpy as np
import openpyxl


NOMBRE_SHEET_TOTAL = "Total"
ESTILO_TABLA = "Table Style Light 9"

def crear_sheets_formato(excel_input_file, output_excel_file = None, column_pivot= None):
    if not isinstance(excel_input_file, list):
        excel_input_file = [excel_input_file]

    if isinstance(excel_input_file[0], pd.DataFrame):
        df = excel_input_file[0]
    
    elif isinstance(excel_input_file[0], str):
        df = pd.read_excel(excel_input_file[0], engine="openpyxl")
    
    else:
        print("tipo de dato no admitido")
        return


def crear_sheets_formato_engine(df, column_pivot , output_excel_file=None):

    columnas = [{"header": column} for column in df.columns]

    pivot_df = df[column_pivot].unique()
    pivot_df = np.sort(pivot_df)

    writer = pd.ExcelWriter(output_excel_file, engine="xlsxwriter")

    tabla_total = df
    tabla_total.to_excel(writer, sheet_name=NOMBRE_SHEET_TOTAL, index=False)
    (max_row, max_col) = tabla_total.shape
    ws_tabla_total = writer.sheets[NOMBRE_SHEET_TOTAL]
    ws_tabla_total.add_table(
        0, 0, max_row, max_col - 1, {"columns": columnas, "style": ESTILO_TABLA}
    )

    for pvt in pivot_df:
        tabla_filtrada = df[df[column_pivot] == pvt]
        (max_row, max_col) = tabla_filtrada.shape
        nombre_sheet = pvt
        tabla_filtrada.to_excel(writer, sheet_name=nombre_sheet, index=False)
        ws_tabla_filtrada = writer.sheets[pvt]
        ws_tabla_filtrada.add_table(
            0, 0, max_row, max_col - 1, {"columns": columnas, "style": ESTILO_TABLA}
        )

    writer.save()

def actualizar_excel_seguimiento(excel_file,  prov_vector):
    wb = openpyxl.load_workbook(excel_file)
    
    with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
        writer.book = wb

        for df, sheet in zip(prov_vector, ['act_data', 'proy_data', 'ref_data']):
            if sheet in wb.sheetnames:
                std = wb[sheet]
                wb.remove(std)

            df.to_excel(writer, sheet_name=sheet, index=False)

        writer.save()
